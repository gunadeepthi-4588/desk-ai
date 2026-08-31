import io
from pypdf import PdfReader
from docx import Document

class TextExtractor:
    @staticmethod
    def extract(file_bytes: bytes, file_type: str) -> list[tuple[int, str]]:
        """
        Extracts text from document bytes based on file type.
        Returns:
            list of tuples: [(page_number, page_text), ...]
            Note: For DOCX and TXT, page_number is defaulted to 1.
        """
        file_type = file_type.lower().strip('.')
        
        if file_type == 'pdf':
            return TextExtractor._extract_pdf(file_bytes)
        elif file_type == 'docx':
            return TextExtractor._extract_docx(file_bytes)
        elif file_type == 'txt':
            return TextExtractor._extract_txt(file_bytes)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> list[tuple[int, str]]:
        pages = []
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            for i, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                pages.append((i, text.strip()))
        except Exception as e:
            raise RuntimeError(f"Failed to extract PDF text: {e}")
        return pages

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> list[tuple[int, str]]:
        try:
            doc = Document(io.BytesIO(file_bytes))
            # Extract paragraph text
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
                    
            # Extract table cell texts
            table_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        table_text.append(" | ".join(row_data))
                        
            if table_text:
                full_text.append("\n=== Document Tables ===\n" + "\n".join(table_text))
                
            combined_text = "\n\n".join(full_text)
            return [(1, combined_text)]
        except Exception as e:
            raise RuntimeError(f"Failed to extract DOCX text: {e}")

    @staticmethod
    def _extract_txt(file_bytes: bytes) -> list[tuple[int, str]]:
        # Try UTF-8 first, fallback to Latin-1
        try:
            text = file_bytes.decode('utf-8')
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode('latin-1')
            except Exception as e:
                raise RuntimeError(f"Failed to decode TXT file: {e}")
        return [(1, text)]
